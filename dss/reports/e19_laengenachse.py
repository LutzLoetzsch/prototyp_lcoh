#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""E19 — Längenachse: Makroergebnis gegen Rechenkern."""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import e19_grenzlaenge as g

SP_LAENGE, SP_PREIS = 7, 19
STARTZEILE = 30


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mappe", required=True)
    ap.add_argument("--kt", type=float, default=500.0)
    ap.add_argument("--lmax", type=float, default=10500.0,
                    help="obere Grenze der Tabelle [m]; die CSV enthält "
                         "weiterhin alle Zeilen")
    ap.add_argument("--out", default="reports/Experimentierlauf_Wuensdorf/"
                                     "E19_toolvergleich")
    a = ap.parse_args()

    from openpyxl import load_workbook
    ws = load_workbook(a.mappe, data_only=True)["FW_Eingaben"]

    g.NORMALISIERT = True

    zeilen = []
    for r in range(STARTZEILE, ws.max_row + 1):
        L = ws.cell(r, SP_LAENGE).value
        p = ws.cell(r, SP_PREIS).value
        if not L or not p:
            continue
        L = float(L)
        werkzeug = float(p) * 1000.0
        kern = g.rechenkern(L, a.kt, 0.0)
        zeilen.append({
            "laenge_m": L,
            "waermeliniendichte": g.WAERMEMENGE / L,
            "werkzeug": werkzeug,
            "rechenkern": kern,
            "differenz": kern - werkzeug,
            "differenz_prozent": (kern / werkzeug - 1.0) * 100.0,
        })

    if not zeilen:
        raise SystemExit("Keine Ergebniszeilen gefunden. Wurde das Makro "
                         "ausgeführt und die Mappe gespeichert?")

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    ziel = out / "e19_laengenachse_exceltool_gegen_rechenkern.csv"
    with ziel.open("w", newline="", encoding="utf-8") as fh:
        s = csv.writer(fh, delimiter=";")
        s.writerow(["trassenlaenge_m", "waermeliniendichte_mwh_pro_m_a",
                    "EXCELTOOL_SAENA_eur_pro_mwh",
                    "RECHENKERN_eur_pro_mwh",
                    "differenz_eur_pro_mwh", "relativer_fehler_prozent"])
        for z in zeilen:
            s.writerow([f"{z['laenge_m']:.0f}",
                        f"{z['waermeliniendichte']:.4f}".replace(".", ","),
                        f"{z['werkzeug']:.3f}".replace(".", ","),
                        f"{z['rechenkern']:.3f}".replace(".", ","),
                        f"{z['differenz']:+.3f}".replace(".", ","),
                        f"{z['differenz_prozent']:+.3f}".replace(".", ",")])

    tab = [z for z in zeilen if z["laenge_m"] <= a.lmax]
    tex = [r"\begin{table}[htbp]", r"\centering", r"\footnotesize",
           r"\renewcommand{\arraystretch}{1.1}%",
           r"\begin{tabular}{lrrrrr}", r"\toprule",
           r"\textbf{Nr.} & \textbf{Trassen\-länge} & "
           r"\textbf{Wärme\-linien\-dichte} & "
           r"\textbf{Planungs\-werkzeug} & \textbf{Rechen\-kern} & "
           r"\textbf{rel. Fehler} \\",
           r"{}& [m] & [MWh/(m$\cdot$a)] & [EUR/MWh] & [EUR/MWh] & [\%] \\",
           r"\midrule"]

    def dz(x: float, n: int) -> str:
        return f"{x:.{n}f}".replace(".", ",")

    for i, z in enumerate(tab):
        farbe = r"\rowcolor{gray!8}" if i % 2 == 0 else ""
        laenge = (f"\\num{{{z['laenge_m']:.0f}}}" if z["laenge_m"] > 999
                  else f"{z['laenge_m']:.0f}")
        fehler = f"{z['differenz_prozent']:+.2f}".replace(".", "{,}")
        tex.append(f"{farbe}P{i+1:02d} & {laenge} & "
                   f"{dz(z['waermeliniendichte'], 3)} & "
                   f"{dz(z['werkzeug'], 2)} & {dz(z['rechenkern'], 2)} & "
                   f"${fehler}$ \\\\")

    tex += [r"\bottomrule", r"\end{tabular}",
            r"\caption[Längenachse: Werkzeug gegen Rechenkern]{Die "
            + str(len(tab)) + r" Stützstellen der Längenachse. "
            r"Quelle: eigene Berechnung}",
            r"\label{tab:e19-laengenachse}", r"\end{table}"]
    tex_ziel = out / "e19_laengenachse.tex"
    tex_ziel.write_text("\n".join(tex) + "\n", encoding="utf-8")
    print(f"CSV: {ziel}\nLaTeX: {tex_ziel}")


if __name__ == "__main__":
    main()
