#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from dss.szenarien.m6 import M6_BASIS, lcoh_m6

import _lauf

TRASSE_M = 5000.0
LEISTUNG_KW = 5000.0
GEBAEUDE = 300
HAUSANSCHLUSS_EUR = 10000.0

ZINS = 0.045
HORIZONT = 20.0
BRENNSTOFF_EUR_MWH = 32.5
NUTZUNGSGRAD = 0.88
VERLUSTANTEIL = 0.163

# Im Werkzeug Anteil der Annuität, im Prototyp Anteil der Investition
WARTUNG_TOOL = 0.015
ETA_TOOL = NUTZUNGSGRAD * (1.0 - VERLUSTANTEIL)

TRASSENKOSTEN = [500.0, 1000.0, 1500.0, 2000.0, 2500.0, 3000.0]  # EUR/m
VOLLBENUTZUNG = [1500, 1850, 2200]  # h/a

ERZEUGER_EUR_KW = 1250.0


def _punkte() -> list[tuple[str, float, int]]:
    return [(f"P{i:02d}", kt, vbh)
            for i, (vbh, kt) in enumerate(
                ((v, k) for v in VOLLBENUTZUNG for k in TRASSENKOSTEN), start=1)]


PUNKTE = _punkte()

STUFEN = {
    "N1": "Hauptvergleich: ohne Gewinnziel, ohne Entfall Ersatzinvestition, "
          "ohne Förderung, ohne CO2, Horizont beidseitig 20 a, Wirkungsgrad "
          "beidseitig zusammengefasst, Wartung in der Werkzeugkonvention "
          "(Anteil der Annuität, im Prototyp umgerechnet)",
}

VARIANTEN = {
    "A": "ohne Hausanschluss (reiner Netz- und Erzeugervergleich)",
    "B": "mit Hausanschluss (Prototyp als eigener Term, Tool als Erzeugeraufschlag)",
}

PROTOKOLL = "e19_excel_protokoll.csv"


def _annuitaetsfaktor(i: float = ZINS, n: float = HORIZONT) -> float:
    return (i * (1 + i) ** n) / ((1 + i) ** n - 1)


def _wartung_invest() -> float:
    return WARTUNG_TOOL * _annuitaetsfaktor()


def waermemenge(vbh: int) -> float:
    return LEISTUNG_KW * vbh / 1000.0


def hausanschluss_eur_pro_mwh_a(vbh: int) -> float:
    return HAUSANSCHLUSS_EUR * GEBAEUDE / waermemenge(vbh)


def prototyp_parameter(kt: float, vbh: int, stufe: str, variante: str) -> dict:
    p = {
        "waermeliniendichte": waermemenge(vbh) / TRASSE_M,
        "netz_invest": kt,
        "capex_bio": ERZEUGER_EUR_KW,
        "vlh_bio": float(vbh),
        "nutzungsgrad": NUTZUNGSGRAD,
        "verlustanteil": VERLUSTANTEIL,
        "brennstoffpreis": BRENNSTOFF_EUR_MWH,
        "wacc": ZINS,
        "betrachtungszeitraum": HORIZONT,
        "wartung_bio": _wartung_invest(),
        "netz_betriebskostenanteil": 0.0,
        "netz_foerderquote": 0.0,
        "foerderquote": 0.0,
        "co2_preis": 0.0,
        # Muss für die Validierung echt größer als 0 sein
        "grundlast_anteil": 1e-9,
        "capex_wp": ERZEUGER_EUR_KW,
        "strompreis": BRENNSTOFF_EUR_MWH,
        "jaz": NUTZUNGSGRAD,
        "vlh_wp": float(vbh),
        "wartung_wp": _wartung_invest(),
        "bedarfsrate": 0.0,
        "hausanschluss": hausanschluss_eur_pro_mwh_a(vbh) if variante == "B" else 0.0,
    }
    return p


def prototyp_zerlegung(p: dict) -> dict:
    voll = lcoh_m6(p)
    ohne_bed = lcoh_m6({**p, "brennstoffpreis": 0.0, "strompreis": 0.0})
    ohne_bet = lcoh_m6({**p, "wartung_bio": 0.0, "wartung_wp": 0.0})
    nur_kap = lcoh_m6({**p, "brennstoffpreis": 0.0, "strompreis": 0.0,
                       "wartung_bio": 0.0, "wartung_wp": 0.0})
    bedarf = voll - ohne_bed
    betrieb = voll - ohne_bet
    kapital = voll - bedarf - betrieb
    rest = kapital - nur_kap
    if abs(rest) > ADD_TOLERANZ:
        raise SystemExit(
            "Zerlegung des Prototyps nicht additiv, Rest "
            f"{rest:.9f} EUR/MWh bei netz_invest={p['netz_invest']}. "
            "Eine der beiden Größen geht nichtlinear ein.")
    return {"gesamt": voll, "kapital": kapital,
            "bedarf": bedarf, "betrieb": betrieb}


def excel_eingaben(kt: float, vbh: int, stufe: str, variante: str) -> dict:
    q_kwh = waermemenge(vbh) * 1000.0
    erz = ERZEUGER_EUR_KW + (HAUSANSCHLUSS_EUR * GEBAEUDE / LEISTUNG_KW
                             if variante == "B" else 0.0)
    return {
        "spez. Kosten Waermetrasse [EUR/m]": kt,
        "Trassenlaenge [m]": TRASSE_M,
        "Waermelieferung p.a. [kWh]": q_kwh,
        "spez. Kosten Waermeerzeugung [EUR/kW]": erz,
        "Normheizlast [kW] (feste Zelle)": LEISTUNG_KW,
        "Wirkungsgrad real": round(ETA_TOOL, 4),
        "Preis Energietraeger real [EUR/kWh]": BRENNSTOFF_EUR_MWH / 1000.0,
        "Wartung/Instandsetzung real [%/a]": WARTUNG_TOOL,
        "Kalkulationszins": ZINS,
        "Kreditlaufzeit [a]": HORIZONT,
        "Gewinnziel": 0.0,
        "Abminderung Entfall Ersatzinvestition [EUR]": 0.0,
    }


def formular(pfad: Path) -> None:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    spalten = ["punkt", "stufe", "variante", "kt_eur_pro_m", "vbh_h_pro_a",
               *excel_eingaben(1000.0, 1850, "N1", "A").keys(),
               "EXCEL_waermepreis_ct_pro_kWh", "datum", "bearbeiter", "bemerkung"]
    with pfad.open("w", newline="", encoding="utf-8") as f:
        s = csv.DictWriter(f, fieldnames=spalten, delimiter=";")
        s.writeheader()
        for stufe in STUFEN:
            for variante in VARIANTEN:
                for kennung, kt, vbh in PUNKTE:
                    zeile = {"punkt": kennung, "stufe": stufe, "variante": variante,
                             "kt_eur_pro_m": kt, "vbh_h_pro_a": vbh,
                             **excel_eingaben(kt, vbh, stufe, variante),
                             "EXCEL_waermepreis_ct_pro_kWh": "",
                             "datum": "", "bearbeiter": "", "bemerkung": ""}
                    s.writerow(zeile)
    print(f"Eingabeprotokoll geschrieben: {pfad}")


MATRIX_STARTZEILE = 30
MATRIX_BLATT = "FW_Eingaben"
SP_TRASSENKOSTEN = 6
SP_WAERMEPREIS = 19
SP_VOLLBENUTZUNG = 21
SP_ERZEUGERKOSTEN = 24
SP_ANNUITAET = 11
SP_WARTUNG = 13
SP_ENERGIE = 15
SP_GESAMT = 17
ADD_TOLERANZ = 1e-6  # EUR/MWh


def lies_excel(pfade: list[Path]) -> dict[tuple[float, int, str], float]:
    from openpyxl import load_workbook

    aufschlag = HAUSANSCHLUSS_EUR * GEBAEUDE / LEISTUNG_KW
    werte: dict[tuple[float, int, str], float] = {}
    for pfad in pfade:
        wb = load_workbook(pfad, data_only=True)
        if MATRIX_BLATT not in wb.sheetnames:
            raise SystemExit(f"{pfad}: Blatt '{MATRIX_BLATT}' fehlt")
        ws = wb[MATRIX_BLATT]
        for r in ws.iter_rows(min_row=MATRIX_STARTZEILE, max_col=24):
            preis = r[SP_WAERMEPREIS - 1].value
            kt_roh = r[SP_TRASSENKOSTEN - 1].value
            # Das Makro schreibt am Ende noch eine Nullzeile
            if preis is None or kt_roh is None:
                continue
            if float(preis) <= 0.0 or float(kt_roh) <= 0.0:
                continue
            kt = float(kt_roh)
            vbh = int(round(float(r[SP_VOLLBENUTZUNG - 1].value)))
            erz = float(r[SP_ERZEUGERKOSTEN - 1].value)
            variante = "B" if erz > ERZEUGER_EUR_KW + aufschlag / 2 else "A"
            mwh = waermemenge(vbh)
            kap = float(r[SP_ANNUITAET - 1].value) / mwh
            bet = float(r[SP_WARTUNG - 1].value) / mwh
            bed = float(r[SP_ENERGIE - 1].value) / mwh
            ges = float(r[SP_GESAMT - 1].value) / mwh
            preis_mwh = float(preis) * 1000.0
            rest = ges - (kap + bed + bet)
            if abs(rest) > ADD_TOLERANZ:
                raise SystemExit(
                    f"{pfad.name} Zeile {r[0].row}: die drei Kostenarten "
                    f"summieren nicht auf die Gesamtkosten, Rest "
                    f"{rest:.9f} EUR/MWh")
            if abs(ges - preis_mwh) > ADD_TOLERANZ:
                raise SystemExit(
                    f"{pfad.name} Zeile {r[0].row}: die Gesamtkosten treffen "
                    f"den Wärmepreis nicht, Differenz "
                    f"{ges - preis_mwh:.9f} EUR/MWh. Möglicher "
                    f"Gewinnaufschlag in Spalte {SP_WAERMEPREIS}.")
            werte[(kt, vbh, variante)] = {
                "gesamt": preis_mwh,
                "kapital": kap, "bedarf": bed, "betrieb": bet,
            }
    if not werte:
        raise SystemExit("Keine Ergebniszeilen gefunden -- wurde das Makro ausgeführt?")
    return werte


def _mlp_werte(saetze: list[dict]) -> list[float | None]:
    try:
        from dss.surrogat.persistenz import lade_modell  # type: ignore
    except Exception:  # pragma: no cover
        return [None] * len(saetze)
    pfad = Path("reports/E07_entscheidung/modelle/e07_bio.pt")
    if not pfad.exists():
        return [None] * len(saetze)
    try:
        modell, sx, sy, meta = lade_modell(pfad)  # pragma: no cover
        import numpy as np
        import torch
        namen = meta["param_namen"]
        X = np.array([[p.get(k, M6_BASIS.get(k, 0.0)) for k in namen] for p in saetze],
                     dtype=float)
        with torch.no_grad():
            y = sy.zurueck(modell(torch.tensor(sx.hin(X), dtype=torch.float32)).numpy())
        return [float(v) for v in np.ravel(y)]
    except Exception:  # pragma: no cover
        return [None] * len(saetze)


def vergleich(pfad: Path, mit_mlp: bool, mappen: list[Path] | None = None) -> dict:
    aus_excel = lies_excel(mappen) if mappen else {}

    if aus_excel:
        zeilen = [{"punkt": k, "stufe": "N1", "variante": v,
                   "kt_eur_pro_m": kt, "vbh_h_pro_a": vbh}
                  for v in VARIANTEN
                  for k, kt, vbh in PUNKTE]
    else:
        if not pfad.exists():
            raise SystemExit(f"Weder Arbeitsmappen übergeben noch Protokoll vorhanden: {pfad}")
        with pfad.open(encoding="utf-8") as f:
            zeilen = list(csv.DictReader(f, delimiter=";"))

    saetze, ergebnis = [], []
    for z in zeilen:
        kt, vbh = float(z["kt_eur_pro_m"]), int(z["vbh_h_pro_a"])
        p = prototyp_parameter(kt, vbh, z["stufe"], z["variante"])
        saetze.append(p)
        if aus_excel:
            excel = aus_excel.get((kt, vbh, z["variante"]))
        else:
            roh = (z.get("EXCEL_waermepreis_ct_pro_kWh") or "").strip().replace(",", ".")
            excel = ({"gesamt": float(roh) * 10.0, "kapital": None,
                      "bedarf": None, "betrieb": None}
                     if roh else None)
        pz = prototyp_zerlegung(p)
        proto = pz["gesamt"]
        ex_ges = None if excel is None else excel["gesamt"]
        zeile = {
            "punkt": z["punkt"], "stufe": z["stufe"], "variante": z["variante"],
            "kt_eur_pro_m": kt, "vbh": vbh,
            "q_l": round(p["waermeliniendichte"], 4),
            "excel_eur_pro_mwh": None if ex_ges is None else round(ex_ges, 3),
            "prototyp_eur_pro_mwh": round(proto, 3),
            "differenz_eur_pro_mwh": None if ex_ges is None else round(proto - ex_ges, 3),
            "differenz_prozent": None if ex_ges is None else round(proto / ex_ges - 1.0, 5),
        }
        for art in ("kapital", "bedarf", "betrieb"):
            pw = pz[art]
            ew = None if excel is None else excel[art]
            zeile["prototyp_" + art] = round(pw, 3)
            zeile["excel_" + art] = None if ew is None else round(ew, 3)
            zeile["differenz_" + art] = None if ew is None else round(pw - ew, 3)
        ergebnis.append(zeile)

    if mit_mlp:
        for e, v in zip(ergebnis, _mlp_werte(saetze)):
            e["surrogat_eur_pro_mwh"] = None if v is None else round(v, 3)
            e["surrogat_minus_analytisch"] = (
                None if v is None else round(v - e["prototyp_eur_pro_mwh"], 4))

    gefuellt = [e for e in ergebnis if e["differenz_prozent"] is not None]
    zus = {}
    if gefuellt:
        d = [abs(e["differenz_prozent"]) for e in gefuellt]
        zus = {"punkte_verglichen": len(gefuellt),
               "mittlere_abweichung": round(sum(d) / len(d), 5),
               "groesste_abweichung": round(max(d), 5),
               "H_B1_erfuellt": max(d) < 0.10}
    return {"szenario": {"trasse_m": TRASSE_M, "leistung_kw": LEISTUNG_KW,
                         "gebaeude": GEBAEUDE, "hausanschluss_eur": HAUSANSCHLUSS_EUR,
                         "zins": ZINS, "horizont_a": HORIZONT},
            "stufen": STUFEN, "varianten": VARIANTEN,
            "zeilen": ergebnis, "zusammenfassung": zus}


def _tabelle(d: dict) -> str:
    z = ["| Punkt | Stufe | Var. | kT [€/m] | Vbh | Q_l | Excel | Prototyp | Δ | Δ % |",
         "|---|---|---|---:|---:|---:|---:|---:|---:|---:|"]
    for e in d["zeilen"]:
        ex = "—" if e["excel_eur_pro_mwh"] is None else f"{e['excel_eur_pro_mwh']:.2f}"
        df = "—" if e["differenz_eur_pro_mwh"] is None else f"{e['differenz_eur_pro_mwh']:+.2f}"
        dp = "—" if e["differenz_prozent"] is None else f"{e['differenz_prozent']:+.2%}"
        z.append(f"| {e['punkt']} | {e['stufe']} | {e['variante']} | {e['kt_eur_pro_m']:.0f} | "
                 f"{e['vbh']} | {e['q_l']:.3f} | {ex} | {e['prototyp_eur_pro_mwh']:.2f} | {df} | {dp} |")
    return "\n".join(z)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--formular", action="store_true", help="Eingabeprotokoll erzeugen")
    p.add_argument("--vergleich", nargs="*", metavar="XLSM", default=None,
                   help="Vergleichstabelle erzeugen; optional die Arbeitsmappen angeben")
    p.add_argument("--mlp", action="store_true", help="Surrogatspalte ergänzen")
    args = p.parse_args()

    ziel = Path(_lauf.reports("E19_toolvergleich"))
    protokoll = ziel / PROTOKOLL

    if args.formular:
        formular(protokoll)
    elif args.vergleich is not None:
        mappen = [Path(x) for x in args.vergleich] or None
        if mappen:
            for m in mappen:
                if not m.exists():
                    raise SystemExit(f"Arbeitsmappe nicht gefunden: {m}")
        d = vergleich(protokoll, args.mlp, mappen)
        ziel.mkdir(parents=True, exist_ok=True)
        (ziel / "e19_vergleich.json").write_text(
            json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
        (ziel / "e19_vergleich.md").write_text(_tabelle(d) + "\n", encoding="utf-8")
        print(f"Geschrieben: {ziel}/e19_vergleich.json und .md")
    else:
        p.print_help()
